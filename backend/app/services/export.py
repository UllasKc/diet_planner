"""Renders a generated diet plan as a DOCX or PDF file for handing to a
client — styled to match the app's green/amber theme rather than a bare
default Word/reportlab look.
"""

from datetime import date
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

# Palette shared with the frontend's green/amber theme.
GREEN_DARK = "16794C"
GREEN = "1F9D63"
GREEN_LIGHT = "E4F6EC"
AMBER_DARK = "D97706"
AMBER = "F59E0B"
AMBER_LIGHT = "FEF3E0"
BLUE = "3B6BF5"
BLUE_LIGHT = "E8EDFE"
PURPLE = "A855F7"
PURPLE_LIGHT = "F4E9FE"
INK = "12261C"
MUTED = "6B7A70"
BORDER = "E5E9E4"
WHITE = "FFFFFF"


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _slug_to_label(value: str) -> str:
    return str(value).replace("_", " ").replace("-", " ").title()


def _ingredient_rows(meal: dict) -> list[tuple[str, str, str, bool]]:
    """Returns (food, quantity_or_choices, calories, is_fixed) tuples."""
    rows = []
    for key, ingredient in meal.get("ingredients", {}).items():
        choices = ingredient.get("choices") or []
        is_fixed = bool(ingredient.get("is_fixed"))
        if choices:
            option_text = " | ".join(
                f"{_slug_to_label(c.get('name', 'Choice'))}: {c.get('quantity', 0)}{c.get('unit', ingredient.get('unit', ''))} ({c.get('calories', 0)} kcal)"
                for c in choices
            )
            rows.append((_slug_to_label(key), option_text, "", is_fixed))
        else:
            rows.append(
                (
                    _slug_to_label(key),
                    f"{ingredient.get('quantity', 0)} {ingredient.get('unit', '')}".strip(),
                    f"{ingredient.get('calories', 0)} kcal",
                    is_fixed,
                )
            )
    return rows


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------


def _shade(element, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    element.append(shd)


def _shade_cell(cell, hex_color: str) -> None:
    _shade(cell._tc.get_or_add_tcPr(), hex_color)


def _shade_paragraph(paragraph, hex_color: str) -> None:
    _shade(paragraph._p.get_or_add_pPr(), hex_color)


def _bar_paragraph(document, text: str, bg: str, fg: str, size: int, space_before: int, space_after: int):
    p = document.add_paragraph()
    _shade_paragraph(p, bg)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(f"  {text}")
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = _rgb(fg)
    return p


def _style_table(table, header_bg: str, header_fg: str, col_widths=None):
    table.style = "Table Grid"
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = 1
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
            if row_idx == 0:
                _shade_cell(cell, header_bg)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = _rgb(header_fg)
    # Table Grid gives black borders by default — recolor to something softer.
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), BORDER)
        borders.append(el)
    tbl_pr.append(borders)


def build_docx(plan: dict) -> BytesIO:
    client = plan["client"]
    nutrition = plan["nutrition"]
    macros = nutrition["macros"]

    document = Document()
    for section in document.sections:
        section.top_margin = Pt(28)
        section.left_margin = Pt(36)
        section.right_margin = Pt(36)

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.font.color.rgb = _rgb(INK)

    # --- Title banner ---
    _bar_paragraph(document, "Personalized Diet Plan", GREEN, WHITE, 24, 0, 2)
    _bar_paragraph(
        document,
        f"{client.get('name') or 'Client'}  •  Generated {date.today().strftime('%d %b %Y')}",
        GREEN,
        GREEN_LIGHT,
        11,
        0,
        14,
    )

    # --- Client info table ---
    info_headers = ["Gender", "Age", "Height", "Weight", "Goal", "Preference"]
    info_values = [
        client["gender"],
        f"{client['age']} yrs",
        f"{client['height_cm']} cm",
        f"{client['weight_kg']} kg",
        client["goal"],
        client.get("food_preference", ""),
    ]
    info_table = document.add_table(rows=2, cols=len(info_headers))
    for idx, header in enumerate(info_headers):
        info_table.rows[0].cells[idx].text = header
        info_table.rows[1].cells[idx].text = str(info_values[idx])
    _style_table(info_table, GREEN_LIGHT, GREEN_DARK)
    document.add_paragraph().paragraph_format.space_after = Pt(4)

    # --- Calorie summary ---
    cal_table = document.add_table(rows=2, cols=3)
    for idx, (label, key) in enumerate([("BMR", "bmr"), ("Maintenance", "maintenance"), ("Target", "target")]):
        cal_table.rows[0].cells[idx].text = label
        cal_table.rows[1].cells[idx].text = f"{nutrition[key]} kcal"
    _style_table(cal_table, GREEN_LIGHT, GREEN_DARK)
    for cell in cal_table.rows[1].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = _rgb(GREEN_DARK)
    document.add_paragraph().paragraph_format.space_after = Pt(4)

    # --- Macro summary (each column colour-coded) ---
    macro_table = document.add_table(rows=2, cols=3)
    macro_table.rows[0].cells[0].text = "Protein"
    macro_table.rows[0].cells[1].text = "Carbs"
    macro_table.rows[0].cells[2].text = "Fats"
    macro_table.rows[1].cells[0].text = f"{macros['protein_g']} g ({macros['protein_calories']} kcal)"
    macro_table.rows[1].cells[1].text = f"{macros['carbs_g']} g ({macros['carbs_calories']} kcal)"
    macro_table.rows[1].cells[2].text = f"{macros['fat_g']} g ({macros['fat_calories']} kcal)"
    table_style = "Table Grid"
    macro_table.style = table_style
    macro_colors = [(BLUE_LIGHT, BLUE), (AMBER_LIGHT, AMBER_DARK), (PURPLE_LIGHT, PURPLE)]
    for col_idx, (bg, fg) in enumerate(macro_colors):
        _shade_cell(macro_table.rows[0].cells[col_idx], bg)
        _shade_cell(macro_table.rows[1].cells[col_idx], bg)
        for row in macro_table.rows:
            for run in row.cells[col_idx].paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = _rgb(fg)
                run.font.size = Pt(10)
    document.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- Meals ---
    last_slot = None
    for meal in plan["meals"]:
        if meal["meal_slot"] != last_slot:
            _bar_paragraph(document, meal["meal_label"].upper(), AMBER, WHITE, 12, 14, 6)
            last_slot = meal["meal_slot"]

        meal_title = document.add_paragraph()
        meal_title.paragraph_format.space_before = Pt(4)
        meal_title.paragraph_format.space_after = Pt(4)
        run = meal_title.add_run(f"{meal.get('meal_name', '')}  ({meal.get('total_calories', 0)} kcal)")
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = _rgb(GREEN_DARK)

        rows = _ingredient_rows(meal)
        if rows:
            table = document.add_table(rows=1 + len(rows), cols=3)
            table.rows[0].cells[0].text = "Food"
            table.rows[0].cells[1].text = "Quantity / Choices"
            table.rows[0].cells[2].text = "Calories"
            for r_idx, (food, qty, kcal, is_fixed) in enumerate(rows, start=1):
                cells = table.rows[r_idx].cells
                cells[0].text = food + ("  [FIXED]" if is_fixed else "")
                cells[1].text = qty
                cells[2].text = kcal
                if is_fixed:
                    for run in cells[0].paragraphs[0].runs:
                        run.font.color.rgb = _rgb(AMBER_DARK)
                if r_idx % 2 == 0:
                    for cell in cells:
                        _shade_cell(cell, "F7FBF9")
            _style_table(table, GREEN_DARK, WHITE)
        document.add_paragraph().paragraph_format.space_after = Pt(2)

    # --- Guidelines ---
    _bar_paragraph(document, "GUIDELINES AND NOTES", GREEN, WHITE, 12, 16, 8)
    for section_data in plan.get("guidelines", []):
        title_p = document.add_paragraph()
        title_p.paragraph_format.space_before = Pt(6)
        title_p.paragraph_format.space_after = Pt(2)
        run = title_p.add_run(section_data["title"])
        run.bold = True
        run.font.color.rgb = _rgb(GREEN_DARK)
        run.font.size = Pt(11)
        for item in section_data["items"]:
            item_p = document.add_paragraph(style="List Bullet")
            item_p.paragraph_format.space_after = Pt(1)
            item_run = item_p.add_run(item)
            item_run.font.size = Pt(10)
            item_run.font.color.rgb = _rgb(INK)

    # --- Footer ---
    footer = document.sections[0].footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run(f"Diet Planner • Generated {date.today().strftime('%d %b %Y')}")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = _rgb(MUTED)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------


def _banner_table(text: str, subtext: str, width: float) -> Table:
    data = [[Paragraph(f"<b>{text}</b>", _pdf_style(18, colors.white))]]
    if subtext:
        data.append([Paragraph(subtext, _pdf_style(10, colors.HexColor(f"#{GREEN_LIGHT}")))])
    table = Table(data, colWidths=[width])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{GREEN}")),
                ("LEFTPADDING", (0, 0), (-1, -1), 14),
                ("TOPPADDING", (0, 0), (-1, 0), 12),
                ("BOTTOMPADDING", (0, -1), (-1, -1), 12),
                ("TOPPADDING", (0, 1), (-1, 1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 2),
            ]
        )
    )
    return table


def _section_bar(text: str, width: float, bg: str = AMBER) -> Table:
    table = Table([[Paragraph(f"<b>{text}</b>", _pdf_style(11, colors.white))]], colWidths=[width])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{bg}")),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def _pdf_style(size: int, color) -> ParagraphStyle:
    return ParagraphStyle(f"s{size}{id(color)}", fontName="Helvetica-Bold", fontSize=size, textColor=color, leading=size + 4)


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
    canvas.drawCentredString(
        doc.pagesize[0] / 2,
        14,
        f"Diet Planner • Generated {date.today().strftime('%d %b %Y')} • Page {doc.page}",
    )
    canvas.restoreState()


def build_pdf(plan: dict) -> BytesIO:
    client = plan["client"]
    nutrition = plan["nutrition"]
    macros = nutrition["macros"]

    styles = getSampleStyleSheet()
    body_style = ParagraphStyle("body", parent=styles["BodyText"], fontSize=8.5, leading=11, textColor=colors.HexColor(f"#{INK}"))
    bold_body = ParagraphStyle("bodyBold", parent=body_style, fontName="Helvetica-Bold")
    fixed_body = ParagraphStyle("bodyFixed", parent=body_style, fontName="Helvetica-Bold", textColor=colors.HexColor(f"#{AMBER_DARK}"))
    meal_title_style = ParagraphStyle("mealTitle", fontName="Helvetica-Bold", fontSize=12, textColor=colors.HexColor(f"#{GREEN_DARK}"), spaceBefore=8, spaceAfter=4)
    section_title_style = ParagraphStyle("sectionTitle", fontName="Helvetica-Bold", fontSize=11, textColor=colors.HexColor(f"#{GREEN_DARK}"), spaceBefore=8, spaceAfter=2)
    bullet_style = ParagraphStyle("bullet", parent=body_style, leftIndent=10, spaceAfter=2)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, topMargin=16 * mm, bottomMargin=16 * mm, leftMargin=16 * mm, rightMargin=16 * mm
    )
    content_width = doc.pagesize[0] - doc.leftMargin - doc.rightMargin

    story = [
        _banner_table(
            "Personalized Diet Plan",
            f"{client.get('name') or 'Client'} &bull; Generated {date.today().strftime('%d %b %Y')}",
            content_width,
        ),
        Spacer(1, 10),
    ]

    info_headers = ["Gender", "Age", "Height", "Weight", "Goal", "Preference"]
    info_values = [
        client["gender"],
        f"{client['age']} yrs",
        f"{client['height_cm']} cm",
        f"{client['weight_kg']} kg",
        client["goal"],
        client.get("food_preference", ""),
    ]
    story.append(_styled_table([info_headers, [str(v) for v in info_values]], header_bg=GREEN_LIGHT, header_fg=GREEN_DARK))
    story.append(Spacer(1, 6))

    calorie_data = [
        ["BMR", "Maintenance", "Target"],
        [f"{nutrition['bmr']} kcal", f"{nutrition['maintenance']} kcal", f"{nutrition['target']} kcal"],
    ]
    story.append(_styled_table(calorie_data, header_bg=GREEN_LIGHT, header_fg=GREEN_DARK, value_fg=GREEN_DARK))
    story.append(Spacer(1, 6))

    macro_data = [
        ["Protein", "Carbs", "Fats"],
        [
            f"{macros['protein_g']} g ({macros['protein_calories']} kcal)",
            f"{macros['carbs_g']} g ({macros['carbs_calories']} kcal)",
            f"{macros['fat_g']} g ({macros['fat_calories']} kcal)",
        ],
    ]
    macro_table = Table(macro_data, colWidths=[content_width / 3] * 3)
    macro_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(f"#{BLUE_LIGHT}")),
                ("BACKGROUND", (1, 0), (1, -1), colors.HexColor(f"#{AMBER_LIGHT}")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor(f"#{PURPLE_LIGHT}")),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor(f"#{BLUE}")),
                ("TEXTCOLOR", (1, 0), (1, -1), colors.HexColor(f"#{AMBER_DARK}")),
                ("TEXTCOLOR", (2, 0), (2, -1), colors.HexColor(f"#{PURPLE}")),
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER}")),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(macro_table)
    story.append(Spacer(1, 10))

    last_slot = None
    for meal in plan["meals"]:
        if meal["meal_slot"] != last_slot:
            story.append(_section_bar(meal["meal_label"].upper(), content_width))
            story.append(Spacer(1, 4))
            last_slot = meal["meal_slot"]

        story.append(Paragraph(f"{meal.get('meal_name', '')} ({meal.get('total_calories', 0)} kcal)", meal_title_style))
        rows = _ingredient_rows(meal)
        if rows:
            table_data = [["Food", "Quantity / Choices", "Calories"]]
            row_styles = []
            for row_idx, (food, qty, kcal, is_fixed) in enumerate(rows, start=1):
                style = fixed_body if is_fixed else body_style
                food_text = food + (" [FIXED]" if is_fixed else "")
                table_data.append([Paragraph(food_text, style), Paragraph(qty, body_style), Paragraph(kcal, body_style)])
                if row_idx % 2 == 0:
                    row_styles.append(("BACKGROUND", (0, row_idx), (-1, row_idx), colors.HexColor("#F7FBF9")))
            table = Table(table_data, colWidths=[content_width * 0.22, content_width * 0.62, content_width * 0.16])
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{GREEN_DARK}")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER}")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                        *row_styles,
                    ]
                )
            )
            story.append(table)
        story.append(Spacer(1, 6))

    story.append(_section_bar("GUIDELINES AND NOTES", content_width, bg=GREEN))
    story.append(Spacer(1, 4))
    for section_data in plan.get("guidelines", []):
        story.append(Paragraph(section_data["title"], section_title_style))
        for item in section_data["items"]:
            story.append(Paragraph(f"&bull; {item}", bullet_style))

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    buffer.seek(0)
    return buffer


def _styled_table(data: list[list[str]], header_bg: str, header_fg: str, value_fg: str | None = None) -> Table:
    table = Table(data)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{header_bg}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(f"#{header_fg}")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER}")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    if value_fg and len(data) > 1:
        style.append(("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor(f"#{value_fg}")))
        style.append(("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"))
    table.setStyle(TableStyle(style))
    return table
